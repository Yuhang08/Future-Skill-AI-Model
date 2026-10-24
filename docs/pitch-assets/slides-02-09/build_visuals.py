from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.graphics.barcode.qr import QrCodeWidget
import json,csv,shutil

ROOT=Path(__file__).resolve().parent
OUT=ROOT if (ROOT/'screens').is_dir() else ROOT/'output'
OUT.mkdir(exist_ok=True)
C={'navy':'0B2038','teal':'087A72','aqua':'36D3BB','ink':'12324C','muted':'49657D','pale':'F3F8FC','line':'C5D9E6','blue':'DCEFFC','white':'FFFFFF','darkpanel':'142E4B','soft':'C1D5E7'}
prs=Presentation();prs.slide_width=Inches(40/3);prs.slide_height=Inches(7.5)
def rgb(c):return RGBColor.from_string(C.get(c,c))
def rect(s,x,y,w,h,fill='white',line='line',radius=True):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(1)
 else:sh.line.fill.background()
 if radius:sh.adjustments[0]=.08
 return sh

def txt(s,text,x,y,w,h,size=20,color='ink',bold=False,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=Inches(.02);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(8)
  if align is not None:p.alignment=align
 return sh

def photo(s,path,x,y,w,h):
 im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
 return s.shapes.add_picture(str(path),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))

def photo_top(s,path,x,y,w,h):
 im=Image.open(path);iw,ih=im.size
 visible=iw*h/w
 if ih>visible:
  p=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h));p.crop_bottom=1-visible/ih
  return p
 return photo(s,path,x,y,w,h)

def line(s,x1,y1,x2,y2,dashed=False,color='teal',arrows=True):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(1.7)
 if dashed:
  from pptx.enum.dml import MSO_LINE_DASH_STYLE
  sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
 if arrows:
  end=OxmlElement('a:tailEnd');end.set('type','triangle');sh.line._get_or_add_ln().append(end)
 return sh

def slide(n,title,subtitle='',dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('navy' if dark else 'pale')
 txt(s,f'FUTURESKILL AI  /  VISUAL FOR SLIDE {n:02d}',.45,.25,11,.3,10,'aqua' if dark else 'teal',True)
 txt(s,title,.45,.74,12.4,.9,30,'white' if dark else 'ink',True)
 if subtitle:txt(s,subtitle,.46,1.57,12.35,.52,15,'soft' if dark else 'muted')
 return s

def footer(s,n,text,dark=False):
 txt(s,text,.46,7.12,12.35,.3,9.5,'soft' if dark else 'muted')

s=slide(2,'Students need a clear next step.','A team-defined problem hypothesis to validate through a student pilot.')
rect(s,.45,2.2,4.55,4.35);photo(s,OUT/'reference-images/classroom-pexels-7092344.jpg',.6,2.37,4.25,3.7)
txt(s,'Illustrative classroom photograph',.68,6.1,4.05,.3,12,'muted')
for y,title,body in [(2.2,'Career direction','Connect interests with roles worth exploring.'),(3.68,'Skill priorities','Identify areas that need further practice.'),(5.16,'Practical action','Turn a chosen direction into weekly activities.')]:
 rect(s,5.3,y,7.55,1.27);rect(s,5.3,y,.07,1.27,'teal',None,False)
 txt(s,title,5.55,y+.12,6.98,.4,21,'ink',True);txt(s,body,5.55,y+.62,6.98,.48,17,'muted')
footer(s,2,'[S2] Pexels photo 7092344. Photo is illustrative; the three needs are not measured research findings.')

s=slide(3,'Connect a career direction to weekly action.','An integrated journey, with direct career browsing available alongside the assessment.')
stages=[('DISCOVER','42-question profile\nor direct browsing'),('EXPLORE','Choose a role\nfrom the catalogue'),('IDENTIFY','Review starter skills\nand self-ratings'),('ACT','Set availability\nand generate tasks'),('TRACK','Completed tasks,\npractice hours, XP')]
for i,(a,b) in enumerate(stages):
 x=.48+i*2.58;rect(s,x,2.6,2.32,1.95)
 txt(s,str(i+1).zfill(2),x+.18,2.78,1,.36,15,'teal',True)
 txt(s,a,x+.18,3.22,2,.4,19,'ink',True);txt(s,b,x+.18,3.78,2,.75,13.5,'muted')
 if i<4:line(s,x+2.35,3.55,x+2.55,3.55)
for x,a,b in [(1.0,'923','catalogued occupations'),(5.0,'42','assessment questions'),(9.0,'3','interface languages')]:
 txt(s,a,x,5.18,3.3,.7,38,'teal',True,PP_ALIGN.CENTER);txt(s,b,x,5.95,3.3,.5,16,'muted',False,PP_ALIGN.CENTER)
footer(s,3,'[S3] Source: project code and supplied catalogue. Skill targets, group starter profiles and time estimates include illustrative assumptions.')

s=slide(4,'Iteration made the experience more consistent.','Actual interface screenshots: an earlier light theme and the current light theme.')
for x,title,name in [(.48,'EARLIER LIGHT THEME','home-before-light.png'),(6.92,'REFINED LIGHT THEME','home-after-light.png')]:
 txt(s,title,x,2.17,5.95,.35,15,'teal',True);rect(s,x,2.64,5.95,3.65);photo_top(s,OUT/'screens'/name,x+.07,2.72,5.81,3.49)
txt(s,'Consistent surfaces  /  readable colours  /  preserved navigation',.6,6.45,12.1,.4,19,'ink',False,PP_ALIGN.CENTER)
footer(s,4,'[S4] Project screenshots; before asset from commit 53cfea7, after capture from 48339bc. This is product iteration, not a user-study result.')

s=slide(5,'Choose a role. Build a plan. Take action.','Illustrative learner: Software Developers • 6 hours/week • an 8-week planning horizon')
for x,title,name in [(.48,'1  CHOOSE A DIRECTION','career-software-developers.png'),(4.78,'2  REVIEW YOUR SKILLS','skills-sample.png'),(9.08,'3  FOLLOW WEEKLY TASKS','weekly-task-detail-sample.png')]:
 txt(s,title,x,2.22,3.78,.65,16,'teal',True);rect(s,x,2.95,3.78,3.65);photo(s,OUT/'screens'/name,x+.06,3.03,3.66,3.49)
footer(s,5,'[S5] Real local browser captures with isolated sample inputs. Full-resolution originals are included. No real learner record or measured outcome is shown.')

s=slide(6,'A connected architecture for the current prototype.','Core career and planning flows use deterministic rules; AI services are a separate integration.')
rect(s,.5,2.65,2.55,1.3);txt(s,'BROWSER',.68,2.84,2.2,.4,18,'ink',True);txt(s,'HTML / CSS / JS',.68,3.34,2.2,.4,13,'muted')
rect(s,4.1,2.65,3.08,1.3,'blue');txt(s,'FASTAPI',4.25,2.84,2.8,.4,18,'ink',True);txt(s,'Routes, sessions and orchestration',4.25,3.34,2.8,.5,13,'muted')
line(s,3.06,3.3,4.06,3.3)
for y,a,b in [(2.03,'CAREER + PLANNING RULES','Matching and starter learning plans'),(3.48,'CSV / O*NET TASK DATA','Occupation profiles and task evidence'),(4.93,'SQLITE RECORDS','Learning, sessions and task progress')]:
 rect(s,9.05,y,3.73,1.13);txt(s,a,9.22,y+.14,3.4,.34,15,'ink',True);txt(s,b,9.22,y+.59,3.4,.4,12.7,'muted');line(s,7.19,3.3,9.02,y+.55)
rect(s,4.1,4.62,3.08,.83);txt(s,'FLASK INTERVIEW MODULE',4.24,4.83,2.8,.4,14,'ink',True);line(s,5.64,3.96,5.64,4.59)
rect(s,4.1,5.99,3.08,.82);txt(s,'OPTIONAL OPENAI SERVICES',4.24,6.17,2.8,.4,13,'ink',True);line(s,5.64,5.46,5.64,5.96,True)
txt(s,'Live provider verification pending',.5,5.98,3.4,.8,15,'muted')
txt(s,'Render: hosting\nGitHub: version control\nReplit / PyCharm: development',.55,4.45,3.2,1.3,14,'muted')
footer(s,6,'[S6] Source: src/main.py, reskilling/storage.py, explanation_provider.py and hr_interview/. Dashed link: AI service integration pending live verification.')

s=slide(7,'Skills development through SDG 4.4.','Keep the platform contribution distinct from exploratory occupation-to-Goal context.',True)
rect(s,.48,2.23,5.35,4.53,'darkpanel','darkpanel');photo(s,OUT/'reference-images/un-sdg-4-official.png',.72,2.48,1.0,1.0)
txt(s,'PLATFORM CONTRIBUTION',1.94,2.54,3.5,.7,18,'white',True)
txt(s,'Identify practice priorities\nFollow learning activities\nTrack completed tasks',.8,3.75,4.7,1.65,20,'soft')
txt(s,'Potential support for skills development.\nImpact is to be evaluated in the pilot.',.8,5.7,4.7,.75,14,'soft')
rect(s,6.12,2.23,6.71,4.53,'white','white');photo(s,OUT/'screens/sdg-task-evidence.png',6.29,2.37,3.29,4.24)
txt(s,'TASK IDs + SOURCES',9.79,2.66,2.8,.7,18,'ink',True)
txt(s,'Solar Energy Systems Engineers\n\nO*NET Task ID 16570\n\nExploratory Target 7.2 hypothesis',9.79,3.6,2.75,2.6,15,'muted')
footer(s,7,'[S7] UN Goal 4 / Target 4.4; O*NET 31.0 Task Statements. Source data does not validate the project’s SDG interpretation or imply UN endorsement.',True)

s=slide(8,'A measurable route to Phase 2.','Separate implemented capabilities, engineering checks and future user outcomes.')
for x,title,body in [(.5,'CURRENT PROTOTYPE','Career exploration\nStarter learning pathways\nSQLite task tracking\nThree-language interface'),(6.85,'PHASE 2 PRIORITIES','Structured student pilot\nValidate role requirements\nVerify live AI services\nSecurity, durable storage and monitoring')]:
 rect(s,x,2.3,5.95,2.68);txt(s,title,x+.25,2.53,5.42,.4,19,'teal',True);txt(s,body,x+.25,3.13,5.38,1.75,18,'ink')
rect(s,.5,5.34,12.3,1.34,'blue','blue')
txt(s,'PROPOSED PILOT: 20 STUDENTS  /  TARGET: AT LEAST 16 SUCCEED (80%)',.78,5.53,11.7,.5,20,'ink',True)
txt(s,'Success: independently create a plan and explain the priority skill and next activity. Baseline: not yet measured.',.78,6.12,11.7,.46,13,'muted')
footer(s,8,'[S8] 80% is the team’s Phase 2 target. A 20-person sample and 16 successes are a proposed test design, not observed results.')

q=QrCodeWidget('https://future-skill-ai-model.onrender.com/');q.qr.make();m=q.qr.modules;cell=12;quiet=4
im=Image.new('RGB',((len(m)+2*quiet)*cell,)*2,'white');d=ImageDraw.Draw(im)
for r,row in enumerate(m):
 for c,value in enumerate(row):
  if value:d.rectangle(((c+quiet)*cell,(r+quiet)*cell,(c+quiet+1)*cell-1,(r+quiet+1)*cell-1),fill='black')
im.save(OUT/'reference-images/live-demo-qr.png')
s=slide(9,'From “What should I do?”','to “What can I learn next?”',True)
txt(s,'Test a clearer plan.\nFind a clearer next step.',.58,2.6,7.7,1.35,29,'white',True)
txt(s,'Proposed ask: student pilot participants\nand career-service feedback.',.6,4.3,7.6,.9,21,'soft')
txt(s,'Francium fluoride\nXiamen University Malaysia',.6,5.8,7.6,.8,18,'aqua',True)
rect(s,9.0,2.25,3.8,3.8,'white','white');photo(s,OUT/'reference-images/live-demo-qr.png',9.17,2.42,3.46,3.46)
txt(s,'SCAN TO TRY THE PROTOTYPE',8.6,6.3,4.4,.4,15,'aqua',True,PP_ALIGN.CENTER)
footer(s,9,'[S9] QR encodes https://future-skill-ai-model.onrender.com/ • URL supplied by the team; live deployment was not verified from this cloud environment.',True)

prs.save(OUT/'FutureSkill_Slides_2-9_Editable_Visuals.pptx')
if Path(__file__).resolve() != (OUT/'build_visuals.py').resolve():
 shutil.copyfile(__file__,OUT/'build_visuals.py')
print('Created editable visual source:',len(prs.slides),'slides')
